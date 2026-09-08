import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_5slide_round1_presentation():
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Visual Theme Palette (Premium Dark Clinical)
    BG_DARK = RGBColor(15, 23, 42)        # Slate 900
    CARD_BG = RGBColor(30, 41, 59)        # Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)    # Slate 700
    PRIMARY_TEAL = RGBColor(20, 184, 166) # Teal 500
    LIGHT_TEAL = RGBColor(94, 234, 212)   # Teal 300
    TEXT_WHITE = RGBColor(248, 250, 252)  # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
    ACCENT_RED = RGBColor(244, 63, 94)    # Rose 500
    ACCENT_BLUE = RGBColor(56, 189, 248)  # Sky 400
    ACCENT_GREEN = RGBColor(34, 197, 94)  # Green 500
    ACCENT_AMBER = RGBColor(245, 158, 11) # Amber 500

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="SIH 2026 | PROBLEM STATEMENT SIH26047 (MEDTECH)"):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.2), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = PRIMARY_TEAL
        line.line.fill.background()

        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.35))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = PRIMARY_TEAL

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.65))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # =========================================================================
    # SLIDE 1: TITLE & EXECUTIVE MISSION
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    # Left accent bar
    glow = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.08), Inches(3.8))
    glow.fill.solid()
    glow.fill.fore_color.rgb = PRIMARY_TEAL
    glow.line.fill.background()

    tbox = s1.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(11.0), Inches(2.3))
    tf = tbox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "MediKiosk"
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    p2 = tf.add_paragraph()
    p2.text = "Autonomous Pre-Consultation Case-Taking & Clinical Intelligence Copilot"
    p2.font.size = Pt(20)
    p2.font.color.rgb = PRIMARY_TEAL
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = "Shifting comprehensive patient history acquisition, document OCR digitization, and emergency triage BEFORE doctor consultation."
    p3.font.size = Pt(13.5)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(8)

    # 3 Info Cards
    meta1 = add_card(s1, Inches(1.1), Inches(4.5), Inches(3.5), Inches(1.8))
    m1_tf = meta1.text_frame
    m1_tf.word_wrap = True
    p = m1_tf.paragraphs[0]
    p.text = "PROBLEM STATEMENT"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_TEAL
    p2 = m1_tf.add_paragraph()
    p2.text = "SIH26047\nPatient Case-Taking Software\nDomain: MedTech / BioTech"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(4)

    meta2 = add_card(s1, Inches(4.9), Inches(4.5), Inches(3.5), Inches(1.8))
    m2_tf = meta2.text_frame
    m2_tf.word_wrap = True
    p = m2_tf.paragraphs[0]
    p.text = "EVALUATION ROUND"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_TEAL
    p2 = m2_tf.add_paragraph()
    p2.text = "Round 1 Pitch\nCore Architecture, Feasibility\n& Phased Upgradation Plan"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(4)

    meta3 = add_card(s1, Inches(8.7), Inches(4.5), Inches(3.5), Inches(1.8))
    m3_tf = meta3.text_frame
    m3_tf.word_wrap = True
    p = m3_tf.paragraphs[0]
    p.text = "TEAM DETAILS"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_TEAL
    p2 = m3_tf.add_paragraph()
    p2.text = "Team: [Insert Team Name]\nLeader: [Insert Leader Name]\nSmart India Hackathon 2026"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(4)

    s1.notes_slide.notes_text_frame.text = (
        "SPEAKER NOTE (30 sec):\n"
        "Good morning respected judges. We present MediKiosk for SIH Problem Statement 26047. "
        "In Indian public hospital OPDs, doctors spend up to 80% of their scarce consultation time on "
        "repetitive history questions and decoding torn paper prescriptions. MediKiosk decouples this intake "
        "process and moves it to an autonomous, multilingual kiosk before the consultation begins."
    )

    # =========================================================================
    # SLIDE 2: THE GROUND REALITY & PROBLEM DEFINITION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "Ground Reality & The Critical Bottlenecks in Indian OPDs")

    # 3 Stat Cards
    c1 = add_card(s2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(1.6))
    c1_tf = c1.text_frame
    c1_tf.word_wrap = True
    p = c1_tf.paragraphs[0]
    p.text = "2 to 5 MINS"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p2 = c1_tf.add_paragraph()
    p2.text = "Average Consultation Time\nPer patient in crowded public hospital OPDs"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = TEXT_WHITE

    c2 = add_card(s2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(1.6))
    c2_tf = c2.text_frame
    c2_tf.word_wrap = True
    p = c2_tf.paragraphs[0]
    p.text = "60% - 80%"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p2 = c2_tf.add_paragraph()
    p2.text = "Time Wasted on History Taking\nRepetitive Q&A + deciphering fragmented paper records"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = TEXT_WHITE

    c3 = add_card(s2, Inches(8.8), Inches(1.6), Inches(3.6), Inches(1.6))
    c3_tf = c3.text_frame
    c3_tf.word_wrap = True
    p = c3_tf.paragraphs[0]
    p.text = "ZERO TRIAGE"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_TEAL
    p2 = c3_tf.add_paragraph()
    p2.text = "In Waiting Room Queues\nLife-threatening emergencies sit undetected in lines"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = TEXT_WHITE

    # Bottom Comprehensive Card
    bot_card = add_card(s2, Inches(0.8), Inches(3.5), Inches(11.6), Inches(3.4))
    b_tf = bot_card.text_frame
    b_tf.word_wrap = True
    bp = b_tf.paragraphs[0]
    bp.text = "FOUR FATAL FAILURES OF EXISTING APPROACHES"
    bp.font.size = Pt(12)
    bp.font.bold = True
    bp.font.color.rgb = PRIMARY_TEAL

    points = [
        ("Fragmented Paper Records: ", "Patients carry torn, faded prescriptions from different clinics; doctors have zero historical continuity."),
        ("Why Mobile Apps Fail: ", "Elderly & rural patients face digital illiteracy, lack of smartphones, and complex English-only portals."),
        ("Physician Burnout & Missed Diagnoses: ", "Examining 100+ patients in a 4-hour shift forces doctors into rapid, superficial questioning."),
        ("Unidentified Red-Flags: ", "Heart attack or stroke symptoms wait in general queues for hours without any early automated triage.")
    ]

    for title, desc in points:
        p = b_tf.add_paragraph()
        p.space_before = Pt(8)
        run1 = p.add_run()
        run1.text = "• " + title
        run1.font.bold = True
        run1.font.size = Pt(12.5)
        run1.font.color.rgb = TEXT_WHITE
        run2 = p.add_run()
        run2.text = desc
        run2.font.size = Pt(12.5)
        run2.font.color.rgb = TEXT_MUTED

    s2.notes_slide.notes_text_frame.text = (
        "SPEAKER NOTE (45 sec):\n"
        "Sir, Indian public healthcare faces a severe volume constraint. A single doctor sees over 100 patients "
        "every morning. In a 3-minute consultation, over 2 minutes are burned repeatedly asking basic questions "
        "and deciphering hand-written prescriptions. Mobile health apps don't work for the majority of these patients "
        "due to age, literacy, and language barriers. Furthermore, life-threatening symptoms sit undetected in "
        "the waiting hall. This calls for an on-premise hardware-software intake kiosk."
    )

    # =========================================================================
    # SLIDE 3: PROPOSED SOLUTION & 4-TIER SYSTEM ARCHITECTURE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Proposed Solution: End-to-End 4-Tier System Architecture")

    tiers = [
        ("TIER 1: IOT EDGE", "ESP32 & Smart RFID", 
         "• Instant contactless patient identification (<1 sec).\n• Physical separation: Card stores ONLY random UID, zero health data (PHI) on card for total privacy.", 
         Inches(0.8)),
        ("TIER 2: PATIENT INTAKE", "Touch & Voice Kiosk", 
         "• Multilingual conversational intake in 13 Indian languages.\n• Native audio prompts + voice input for low-literacy users.\n• On-kiosk HD document scanner for old prescriptions.", 
         Inches(3.8)),
        ("TIER 3: CORE SAFETY BUS", "Realtime Triage Engine", 
         "• Deterministic emergency triage rules (Chest pain, dyspnea).\n• Zero-LLM latency: immediate audio-visual alert to hospital nursing queue.\n• WebSockets + Prisma ORM + PostgreSQL.", 
         Inches(6.8)),
        ("TIER 4: CLINICAL SUITE", "Doctor 360 & Copilot", 
         "• Live priority-sorted patient queue with risk badges.\n• Evidence-linked SOAP case summary ready before entry.\n• Interactive clinical Q&A copilot + ABDM/FHIR readiness.", 
         Inches(9.8))
    ]

    for tag, title, details, left_pos in tiers:
        card = add_card(s3, left_pos, Inches(1.6), Inches(2.8), Inches(5.2))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_TEAL

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = details
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(10)

        p4 = tf.add_paragraph()
        p4.space_before = Pt(18)
        p4.text = "Status: Architecture Designed & Verified"
        p4.font.size = Pt(8.5)
        p4.font.color.rgb = ACCENT_GREEN

    s3.notes_slide.notes_text_frame.text = (
        "SPEAKER NOTE (45 sec):\n"
        "Our architecture decouples data collection from the consultation room into 4 seamless tiers. "
        "At the edge, a patient taps an RFID card costing less than 15 rupees. The kiosk guides them in their "
        "mother tongue using voice and large touch icons, and scans their previous prescriptions. "
        "While they walk to the doctor's room, our deterministic engine checks for red flags, and the physician "
        "receives an evidence-linked case sheet instantly."
    )

    # =========================================================================
    # SLIDE 4: CORE INNOVATIONS & AI SAFETY (JUDGE HOOK)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Core Technical USPs & Clinical AI Safety Guardrails")

    usps = [
        ("🛡️ Traceable Evidence Source Citations", 
         "Every clinical statement generated by the AI copilot contains clickable citation chips ([Answer #12], [Doc #03]). The physician can click to trace back directly to raw voice transcripts or scanned prescription images, ensuring strict ground-truth auditability.", Inches(0.8), Inches(1.6)),
        ("⚡ Deterministic Emergency Red-Flag Triage", 
         "Life-threatening symptoms (chest pain radiating to arm, acute breathlessness, sudden syncope) bypass probabilistic LLMs completely. A hardcoded, deterministic rule engine triggers instant hospital alarms in under 1 second.", Inches(6.8), Inches(1.6)),
        ("🌿 Dual Allopathic & AYUSH Clinical Engine", 
         "Recognizing India's healthcare diversity, the kiosk supports dynamic branching into AYUSH clinical case-taking (Prakriti, Vikriti, Agni, Dhatu, Sara, Sattva) alongside standard allopathic clinical trees.", Inches(0.8), Inches(4.3)),
        ("📄 Intelligent Rx Layout Analysis & OCR", 
         "Specialized document OCR pipeline detects medical entities (medications, dosages, frequencies, lab values) from crumpled paper records and organizes them into chronological clinical timelines.", Inches(6.8), Inches(4.3))
    ]

    for title, desc, left_pos, top_pos in usps:
        card = add_card(s4, left_pos, top_pos, Inches(5.6), Inches(2.5))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = LIGHT_TEAL

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(8)

    s4.notes_slide.notes_text_frame.text = (
        "SPEAKER NOTE (45 sec):\n"
        "Judges often question whether AI can be trusted in medical environments. Our answer is absolute safety guardrails. "
        "First, emergency red-flags do not rely on LLMs; they are hardcoded deterministic rules. Second, our copilot "
        "enforces strict evidence citations: every claim has a clickable link back to an answer or prescription "
        "box, leaving the doctor in full audit control. Third, we natively support both Allopathic and AYUSH clinical models."
    )

    # =========================================================================
    # SLIDE 5: UPGRADATION & FUTURE INNOVATION ROADMAP
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Phased Upgradation & Future Expansion Roadmap", "FUTURE UPGRADATION & ROADMAP")

    upgrades = [
        ("PHASE 1 (CURRENT ROUND)", "Clinical Foundation & Architecture", 
         ["• SIH26047 problem validation & OPD survey.",
          "• Monorepo architecture & strictly typed schemas.",
          "• Clinical decision trees & deterministic red flags.",
          "• Focus: Design feasibility & safety validation."],
         ACCENT_GREEN, Inches(0.8)),
        ("PHASE 2 (ROUNDS 2-3)", "Software Integration & AI Pipeline", 
         ["• Voice-enabled 13-language patient kiosk UI.",
          "• Rx & Lab report OCR layout analysis pipeline.",
          "• Physician 360 Copilot with evidence citations.",
          "• Multi-facility hospital queue monitoring."],
         ACCENT_BLUE, Inches(3.8)),
        ("PHASE 3 (FINALE / ROUND 4)", "Hardware Fleet & Edge Testing", 
         ["• Physical ESP32 RFID hardware demonstration.",
          "• Central Command Center with AI CER governance.",
          "• Realtime WebSocket latency stress tests (<100ms).",
          "• Full end-to-end patient journey verification."],
         LIGHT_TEAL, Inches(6.8)),
        ("PHASE 4 (FUTURE UPGRADES)", "IoT Vitals Pod & National ABHA", 
         ["• IoT Diagnostic Vitals Pod: Automated BP cuff, SpO2 & IR Thermometer integrated into kiosk.",
          "• Drug-Allergy & Contradiction AI Detection engine.",
          "• National ABHA / ABDM health locker auto-sync.",
          "• District Epidemiological Outbreak AI Heatmaps."],
         ACCENT_AMBER, Inches(9.8))
    ]

    for tag, title, items, color, left_pos in upgrades:
        card = add_card(s5, left_pos, Inches(1.6), Inches(2.8), Inches(5.2))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(13.5)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(4)

        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = item
            p_item.font.size = Pt(10)
            p_item.font.color.rgb = TEXT_MUTED
            p_item.space_before = Pt(7)

    s5.notes_slide.notes_text_frame.text = (
        "SPEAKER NOTE (45 sec):\n"
        "To conclude with our upgradation roadmap: In Round 1, we validated the architecture and clinical logic. "
        "In Rounds 2 and 3, we build out the multilingual voice interface, OCR pipeline, and doctor copilot. "
        "In Round 4, we showcase the live physical ESP32 RFID hardware. "
        "Looking ahead to post-hackathon national deployment, our Phase 4 upgrades integrate an IoT diagnostic "
        "vitals pod—automatically measuring blood pressure, SpO2, and temperature at the kiosk—along with full "
        "ABDM/ABHA health locker synchronization and district-level disease outbreak analytics. "
        "Thank you, we are now ready for your questions!"
    )

    output_path = "c:\\Users\\vishn\\SIH-PS-26047\\MediKiosk_Round1_Presentation.pptx"
    prs.save(output_path)
    print(f"5-Slide Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_5slide_round1_presentation()
